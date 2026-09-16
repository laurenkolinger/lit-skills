"""End-to-end tests: a real ingest run over real PDFs, checked by an independent count."""

import csv
import os
import shutil

from conftest import make_paper
from litkit import index, ingest, sheet

PAPERS = [
    ("reef_cover.pdf", "Coral cover decline across the Caribbean shelf", "Smith, Tyler B.; Brandt, Marilyn E.", 2016, "10.1234/reef.001"),
    ("auv_survey.pdf", "Autonomous underwater vehicle surveys of mesophotic reefs", "Olinger, Lauren K.; Nemeth, Richard S.", 2023, "10.1234/auv.002"),
    ("fish_spawning.pdf", "Red hind spawning aggregation recovery after protection", "Nemeth, Richard S.", 2005, "10.1234/fsa.003"),
]


def stage_papers(ingest_folder, papers=PAPERS):
    """Write each paper into an ingest folder and return how many were written."""
    for name, title, authors, year, doi in papers:
        make_paper(
            os.path.join(ingest_folder, name),
            title,
            authors,
            year,
            doi,
            f"This study examined {title.lower()} using field surveys at multiple sites and "
            "reports the resulting change over the study period.",
        )
    return len(papers)


class TestFullRun:
    def test_three_papers_produce_three_rows_three_files_and_three_archives(self, tmp_path):
        root = str(tmp_path / "Lit")
        paths = index.ensure_layout(root)
        expected = stage_papers(paths["ingest"])

        result = ingest.run(root, source="test batch", use_network=False)

        assert result["added"] == expected
        assert result["rejected"] == 0
        assert result["duplicates"] == 0

        # Independent counts, taken from the filesystem rather than the return value.
        pdf_count = len([f for f in os.listdir(paths["pdfs"]) if f.endswith(".pdf")])
        archived = len([f for f in os.listdir(paths["processed"]) if f.endswith(".pdf")])
        with open(paths["index"], newline="", encoding="utf-8") as handle:
            csv_rows = list(csv.DictReader(handle))

        assert pdf_count == expected
        assert archived == expected
        assert len(csv_rows) == expected
        assert os.listdir(paths["ingest"]) == ["_processed"]

    def test_every_row_points_at_a_file_that_exists(self, tmp_path):
        root = str(tmp_path / "Lit")
        paths = index.ensure_layout(root)
        stage_papers(paths["ingest"])
        ingest.run(root, use_network=False)

        for row in index.read_index(paths["index"]):
            assert os.path.exists(os.path.join(paths["pdfs"], row["filename"]))

    def test_audit_agrees_after_a_clean_run(self, tmp_path):
        root = str(tmp_path / "Lit")
        paths = index.ensure_layout(root)
        stage_papers(paths["ingest"])
        ingest.run(root, use_network=False)

        report = index.audit(root)
        assert report["missing_files"] == []
        assert report["unindexed_files"] == []
        assert report["rows"] == report["files"] == len(PAPERS)

    def test_metadata_extracted_from_the_files_reaches_the_rows(self, tmp_path):
        root = str(tmp_path / "Lit")
        paths = index.ensure_layout(root)
        stage_papers(paths["ingest"])
        ingest.run(root, use_network=False)

        rows = {row["doi"]: row for row in index.read_index(paths["index"])}
        assert "10.1234/auv.002" in rows
        auv = rows["10.1234/auv.002"]
        assert "Autonomous underwater vehicle" in auv["title"]
        assert auv["year"] == "2023"
        assert auv["key"].startswith("Olinger_2023_")

    def test_the_run_writes_a_log_naming_every_added_paper(self, tmp_path):
        root = str(tmp_path / "Lit")
        paths = index.ensure_layout(root)
        stage_papers(paths["ingest"])
        result = ingest.run(root, use_network=False)

        assert os.path.exists(result["log_path"])
        log = open(result["log_path"], encoding="utf-8").read()
        assert log.count("ADD ") == len(PAPERS)

    def test_the_workbook_is_rebuilt_with_a_row_per_paper(self, tmp_path):
        from openpyxl import load_workbook

        root = str(tmp_path / "Lit")
        paths = index.ensure_layout(root)
        stage_papers(paths["ingest"])
        ingest.run(root, use_network=False)

        workbook_path = os.path.join(root, ingest.WORKBOOK_FILENAME)
        assert os.path.exists(workbook_path)
        loaded = load_workbook(workbook_path)[sheet.SHEET_TITLE]
        assert loaded.max_row == len(PAPERS) + 1
        headers = [c.value for c in loaded[1]]
        assert headers[:-1] == index.COLUMNS and headers[-1] == sheet.MATCH_COLUMN
        assert loaded.freeze_panes == "C2"
        assert loaded.auto_filter.ref is not None

    def test_dry_run_changes_nothing_on_disk(self, tmp_path):
        root = str(tmp_path / "Lit")
        paths = index.ensure_layout(root)
        expected = stage_papers(paths["ingest"])

        result = ingest.run(root, dry_run=True, use_network=False)

        assert result["would_add"] == expected
        assert result["added"] == 0
        assert os.listdir(paths["pdfs"]) == []
        assert not os.path.exists(paths["index"])
        staged = [f for f in os.listdir(paths["ingest"]) if f.endswith(".pdf")]
        assert len(staged) == expected

    def test_a_second_batch_appends_rather_than_replacing(self, tmp_path):
        root = str(tmp_path / "Lit")
        paths = index.ensure_layout(root)
        stage_papers(paths["ingest"], PAPERS[:2])
        first = ingest.run(root, use_network=False)

        stage_papers(paths["ingest"], PAPERS[2:])
        second = ingest.run(root, use_network=False)

        assert first["rows_total"] == 2
        assert second["added"] == 1
        assert second["rows_total"] == 3


class TestCommandLine:
    def test_main_reports_and_exits_zero(self, tmp_path, capsys):
        root = str(tmp_path / "Lit")
        paths = index.ensure_layout(root)
        stage_papers(paths["ingest"])

        code = ingest.main([root, "--no-network", "--source", "cli test"])
        output = capsys.readouterr().out

        assert code == 0
        assert "added 3" in output

    def test_main_on_a_missing_root_creates_it_and_reports_zero(self, tmp_path, capsys):
        code = ingest.main([str(tmp_path / "brand_new"), "--no-network"])
        assert code == 0
        assert "added 0" in capsys.readouterr().out


class TestBulkImport:
    def test_import_from_leaves_the_originals_in_place(self, tmp_path):
        source_folder = tmp_path / "elsewhere"
        source_folder.mkdir()
        expected = stage_papers(str(source_folder))
        root = str(tmp_path / "Lit")

        result = ingest.import_from(root, [str(source_folder)], use_network=False)

        assert result["added"] == expected
        assert len([f for f in os.listdir(source_folder) if f.endswith(".pdf")]) == expected
        paths = index.library_paths(root)
        assert len([f for f in os.listdir(paths["pdfs"]) if f.endswith(".pdf")]) == expected
        assert os.listdir(paths["processed"]) == []

    def test_import_from_several_folders_dedupes_across_them(self, tmp_path):
        first, second = tmp_path / "one", tmp_path / "two"
        first.mkdir()
        second.mkdir()
        stage_papers(str(first))
        stage_papers(str(second))  # the same three papers again
        root = str(tmp_path / "Lit")

        result = ingest.import_from(root, [str(first), str(second)], use_network=False)

        assert result["added"] == len(PAPERS)
        assert result["duplicates"] == len(PAPERS)

    def test_auto_source_records_the_originating_folder(self, tmp_path):
        folder = tmp_path / "grouperlit"
        folder.mkdir()
        stage_papers(str(folder), PAPERS[:1])
        root = str(tmp_path / "Lit")

        ingest.import_from(root, [str(folder)], source=ingest.AUTO_SOURCE, use_network=False)

        rows = index.read_index(index.library_paths(root)["index"])
        assert rows[0]["source"] == "grouperlit"

    def test_explicit_source_overrides_auto_detection(self, tmp_path):
        folder = tmp_path / "grouperlit"
        folder.mkdir()
        stage_papers(str(folder), PAPERS[:1])
        root = str(tmp_path / "Lit")

        ingest.import_from(root, [str(folder)], source="hand picked", use_network=False)

        rows = index.read_index(index.library_paths(root)["index"])
        assert rows[0]["source"] == "hand picked"

    def test_recursive_import_reaches_subfolders(self, tmp_path):
        top = tmp_path / "top"
        nested = top / "nested"
        nested.mkdir(parents=True)
        stage_papers(str(nested), PAPERS[:2])
        root = str(tmp_path / "Lit")

        flat = ingest.import_from(root, [str(top)], use_network=False)
        assert flat["added"] == 0

        deep = ingest.import_from(root, [str(top)], recursive=True, use_network=False)
        assert deep["added"] == 2

    def test_missing_source_folder_is_skipped_not_fatal(self, tmp_path):
        root = str(tmp_path / "Lit")
        result = ingest.import_from(root, [str(tmp_path / "nope")], use_network=False)
        assert result["added"] == 0
