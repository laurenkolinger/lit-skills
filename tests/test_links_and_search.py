"""Tests for the link column, the derived tags cell, and the Search and Tags tabs."""

import json
import os

import pytest
from openpyxl import load_workbook

from conftest import make_paper
from litkit import index, links, sheet


class TestLinkBuilding:
    def test_a_known_id_becomes_a_direct_file_link(self):
        assert links.file_link("1AbC") == "https://drive.google.com/file/d/1AbC/view"

    @pytest.mark.parametrize("empty", [None, "", "   "])
    def test_no_id_yields_no_direct_link(self, empty):
        assert links.file_link(empty) == ""

    def test_the_fallback_searches_drive_for_the_filename(self):
        url = links.search_link("Smith_2016_CoralCover.pdf")
        assert url.startswith("https://drive.google.com/drive/search?q=")
        assert "Smith_2016_CoralCover" in url
        assert ".pdf" not in url

    def test_a_filename_with_spaces_is_url_encoded(self):
        assert " " not in links.search_link("A paper with spaces.pdf")

    @pytest.mark.parametrize("empty", [None, "", "   "])
    def test_no_filename_yields_no_fallback(self, empty):
        assert links.search_link(empty) == ""

    def test_a_missing_ids_file_is_not_fatal(self, tmp_path):
        assert links.load_ids(str(tmp_path / "nope.json")) == {}

    def test_a_corrupt_ids_file_is_not_fatal(self, tmp_path):
        bad = tmp_path / "bad.json"
        bad.write_text("{not json")
        assert links.load_ids(str(bad)) == {}

    def test_an_ids_file_holding_a_list_is_rejected(self, tmp_path):
        bad = tmp_path / "list.json"
        bad.write_text('["a", "b"]')
        assert links.load_ids(str(bad)) == {}


@pytest.fixture
def library(tmp_path):
    root = str(tmp_path / "Lit")
    paths = index.ensure_layout(root)
    for name in ("A_2016_X.pdf", "B_2017_Y.pdf"):
        make_paper(os.path.join(paths["pdfs"], name), "T", "A", 2016, "10.1/x", "abs")
    index.write_index(paths["index"], [
        {"key": "A_2016_X", "filename": "A_2016_X.pdf", "title": "Reef cover over time",
         "topic_tags": "coral reef, monitoring", "method_tags": "AUV", "region_tags": "USVI",
         "vicar_relevance": "reef research", "citations": "40", "year": "2016"},
        {"key": "B_2017_Y", "filename": "B_2017_Y.pdf", "title": "Grouper spawning",
         "topic_tags": "fish spawning aggregation", "region_tags": "USVI",
         "vicar_relevance": "reef research", "citations": "9", "year": "2017"},
    ])
    return root, paths


class TestApplyLinks:
    def test_known_ids_become_direct_links(self, tmp_path, library):
        root, paths = library
        ids = tmp_path / "ids.json"
        ids.write_text(json.dumps({"A_2016_X.pdf": "ID1", "B_2017_Y.pdf": "ID2"}))
        result = links.apply_links(root, ids_path=str(ids))
        assert result["direct"] == 2 and result["fallback"] == 0
        rows = index.read_index(paths["index"])
        assert all(r["link"].startswith("https://drive.google.com/file/d/") for r in rows)

    def test_an_unknown_id_still_gets_a_working_link(self, tmp_path, library):
        root, paths = library
        ids = tmp_path / "ids.json"
        ids.write_text(json.dumps({"A_2016_X.pdf": "ID1"}))
        result = links.apply_links(root, ids_path=str(ids))
        assert result["direct"] == 1 and result["fallback"] == 1
        rows = {r["key"]: r for r in index.read_index(paths["index"])}
        assert rows["B_2017_Y"]["link"]
        assert "search" in rows["B_2017_Y"]["link"]

    def test_every_row_ends_up_with_a_link(self, library):
        root, paths = library
        links.apply_links(root)
        assert all(r["link"] for r in index.read_index(paths["index"]))

    def test_dry_run_changes_nothing(self, library):
        root, paths = library
        before = open(paths["index"], "rb").read()
        links.apply_links(root, dry_run=True)
        assert open(paths["index"], "rb").read() == before


class TestTagsAllCell:
    def test_it_is_derived_on_write_and_wrapped_in_pipes(self, library):
        root, paths = library
        row = {r["key"]: r for r in index.read_index(paths["index"])}["A_2016_X"]
        assert row["tags_all"] == "|coral reef|monitoring|AUV|USVI|reef research|"

    def test_a_whole_tag_matches_but_a_fragment_does_not(self, library):
        root, paths = library
        cell = {r["key"]: r for r in index.read_index(paths["index"])}["A_2016_X"]["tags_all"]
        assert "|AUV|" in cell          # the tag itself
        assert "|AU|" not in cell       # a fragment of it must not match

    def test_it_cannot_drift_from_the_facets(self, library):
        root, paths = library
        rows = index.read_index(paths["index"])
        for r in rows:
            if r["key"] == "A_2016_X":
                r["topic_tags"] = "bleaching"
                r["tags_all"] = "|stale|nonsense|"
        index.write_index(paths["index"], rows)
        refreshed = {r["key"]: r for r in index.read_index(paths["index"])}["A_2016_X"]
        assert "stale" not in refreshed["tags_all"]
        assert "|bleaching|" in refreshed["tags_all"]

    def test_a_row_with_no_tags_has_an_empty_cell(self):
        assert index.build_tags_all({"topic_tags": "", "method_tags": None}) == ""

    def test_a_tag_repeated_across_facets_appears_once(self):
        cell = index.build_tags_all({"topic_tags": "reef fish", "taxa_tags": "reef fish"})
        assert cell == "|reef fish|"


class TestWorkbookTabs:
    def build(self, tmp_path, library):
        root, paths = library
        links.apply_links(root)
        out = os.path.join(root, "lit_index.xlsx")
        return load_workbook(out)

    def test_the_search_tab_is_first_so_it_is_what_opens(self, tmp_path, library):
        assert self.build(tmp_path, library).sheetnames[0] == sheet.SEARCH_TITLE

    def test_all_three_tabs_exist(self, tmp_path, library):
        wb = self.build(tmp_path, library)
        assert set(wb.sheetnames) == {sheet.SEARCH_TITLE, sheet.SHEET_TITLE, sheet.TAGS_TITLE}

    def test_the_tags_tab_counts_each_tag(self, tmp_path, library):
        tab = self.build(tmp_path, library)[sheet.TAGS_TITLE]
        found = {r[1]: r[2] for r in tab.iter_rows(min_row=5, values_only=True) if r[1]}
        assert found["USVI"] == 2 and found["AUV"] == 1 and found["reef research"] == 2

    def test_the_link_cell_is_a_real_hyperlink(self, tmp_path, library):
        cell = self.build(tmp_path, library)[sheet.SHEET_TITLE].cell(row=2, column=1)
        assert cell.value == sheet.LINK_TEXT and cell.hyperlink is not None

    def test_the_index_rows_are_still_short(self, tmp_path, library):
        tab = self.build(tmp_path, library)[sheet.SHEET_TITLE]
        assert {tab.row_dimensions[r].height for r in range(2, tab.max_row + 1)} == {sheet.DATA_ROW_HEIGHT}


class TestNoDynamicArrayFunctions:
    """The lesson from 2026-09-16.

    Google Sheets opened the workbook, could not represent FILTER and SORT in the xlsx format,
    replaced the formula with __xludf.DUMMYFUNCTION and saved that over the file. Nothing in
    this workbook may use a function the format cannot carry.
    """

    BANNED = ("FILTER(", "SORT(", "SORTN(", "QUERY(", "UNIQUE(", "SEQUENCE(",
              "XLOOKUP(", "LET(", "LAMBDA(", "TEXTJOIN(", "IFS(", "ARRAYFORMULA(")

    def formulas(self, tmp_path, library):
        root, _ = library
        links.apply_links(root)
        wb = load_workbook(os.path.join(root, "lit_index.xlsx"))
        out = []
        for name in wb.sheetnames:
            for row in wb[name].iter_rows():
                for cell in row:
                    if isinstance(cell.value, str) and cell.value.startswith("="):
                        out.append((name, cell.coordinate, cell.value))
        return out

    def test_no_banned_function_appears_anywhere(self, tmp_path, library):
        offenders = [
            (s, c, f) for s, c, f in self.formulas(tmp_path, library)
            for b in self.BANNED if b in f.upper()
        ]
        assert offenders == [], f"dynamic-array functions found: {offenders[:3]}"

    def test_no_array_literal_appears_anywhere(self, tmp_path, library):
        assert not [f for _, _, f in self.formulas(tmp_path, library) if "{" in f]

    def test_the_workbook_still_has_formulas(self, tmp_path, library):
        """Guard against the ban being satisfied by having no formulas at all."""
        assert len(self.formulas(tmp_path, library)) > 1


class TestMatchColumn:
    def build(self, tmp_path, library):
        root, _ = library
        links.apply_links(root)
        return load_workbook(os.path.join(root, "lit_index.xlsx"))

    def test_the_index_has_a_match_column_after_the_index_columns(self, tmp_path, library):
        from litkit import index as index_module

        tab = self.build(tmp_path, library)[sheet.SHEET_TITLE]
        headers = [c.value for c in tab[1]]
        assert headers[:-1] == index_module.COLUMNS
        assert headers[-1] == sheet.MATCH_COLUMN

    def test_every_data_row_carries_a_match_formula(self, tmp_path, library):
        tab = self.build(tmp_path, library)[sheet.SHEET_TITLE]
        column = len([c for c in tab[1]])
        for line in range(2, tab.max_row + 1):
            value = tab.cell(row=line, column=column).value
            assert isinstance(value, str) and value.startswith("=AND(")

    def test_the_match_formula_reads_the_search_boxes(self, tmp_path, library):
        tab = self.build(tmp_path, library)[sheet.SHEET_TITLE]
        formula = tab.cell(row=2, column=len([c for c in tab[1]])).value
        for box in ("$B$3", "$B$4", "$B$5", "$B$6"):
            assert f"Search!{box}" in formula

    def test_an_empty_box_does_not_exclude_a_row(self, tmp_path, library):
        tab = self.build(tmp_path, library)[sheet.SHEET_TITLE]
        formula = tab.cell(row=2, column=len([c for c in tab[1]])).value
        assert formula.count('="",TRUE') >= 3

    def test_tags_are_matched_whole_not_as_fragments(self, tmp_path, library):
        tab = self.build(tmp_path, library)[sheet.SHEET_TITLE]
        formula = tab.cell(row=2, column=len([c for c in tab[1]])).value
        assert '"|"&Search!$B$3&"|"' in formula

    def test_free_text_searches_title_and_summary(self, tmp_path, library):
        tab = self.build(tmp_path, library)[sheet.SHEET_TITLE]
        formula = tab.cell(row=2, column=len([c for c in tab[1]])).value
        assert formula.count("Search!$B$6") == 3  # the empty test plus title and summary

    def test_the_match_column_is_not_in_the_csv(self, tmp_path, library):
        from litkit import index as index_module

        root, paths = library
        assert sheet.MATCH_COLUMN not in index_module.COLUMNS
        assert sheet.MATCH_COLUMN not in index_module.read_index(paths["index"])[0]

    def test_the_search_tab_counts_the_matches(self, tmp_path, library):
        tab = self.build(tmp_path, library)[sheet.SEARCH_TITLE]
        assert str(tab["B16"].value).startswith("=COUNTIF(")

    def test_the_search_tab_tells_the_reader_what_to_do(self, tmp_path, library):
        tab = self.build(tmp_path, library)[sheet.SEARCH_TITLE]
        body = " ".join(str(c.value) for row in tab.iter_rows() for c in row if c.value)
        assert "filter arrow" in body and sheet.MATCH_COLUMN in body
