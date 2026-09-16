"""Regression tests pinning behaviour that must not drift.

Each test here names a specific promise made to the lab: reruns are safe, the index stays
in step with the folder, and the column set never changes silently.
"""

import os
import shutil

import pytest

from conftest import make_paper
from litkit import index, ingest, naming, sheet


def stage(ingest_folder, name, title, doi, authors="Smith, Tyler B.", year=2016):
    """Write one journal-style PDF into an ingest folder."""
    return make_paper(
        os.path.join(ingest_folder, name), title, authors, year, doi,
        "An abstract describing the surveys, the sites, and the change observed.",
    )


class TestRerunsAreSafe:
    def test_reingesting_the_same_paper_adds_no_row(self, tmp_path):
        root = str(tmp_path / "Lit")
        paths = index.ensure_layout(root)
        stage(paths["ingest"], "a.pdf", "Coral reef cover through time", "10.1/a")
        ingest.run(root, use_network=False)

        # The same paper arrives again, under a different filename.
        stage(paths["ingest"], "a_copy.pdf", "Coral reef cover through time", "10.1/a")
        second = ingest.run(root, use_network=False)

        assert second["added"] == 0
        assert second["duplicates"] == 1
        assert second["rows_total"] == 1

    def test_a_rerun_leaves_the_index_byte_identical(self, tmp_path):
        root = str(tmp_path / "Lit")
        paths = index.ensure_layout(root)
        stage(paths["ingest"], "a.pdf", "Coral reef cover through time", "10.1/a")
        ingest.run(root, use_network=False)
        before = open(paths["index"], "rb").read()

        stage(paths["ingest"], "a_again.pdf", "Coral reef cover through time", "10.1/a")
        ingest.run(root, use_network=False)

        assert open(paths["index"], "rb").read() == before

    def test_a_rerun_does_not_add_a_second_pdf_to_the_folder(self, tmp_path):
        root = str(tmp_path / "Lit")
        paths = index.ensure_layout(root)
        stage(paths["ingest"], "a.pdf", "Coral reef cover through time", "10.1/a")
        ingest.run(root, use_network=False)

        stage(paths["ingest"], "a_again.pdf", "Coral reef cover through time", "10.1/a")
        ingest.run(root, use_network=False)

        assert len([f for f in os.listdir(paths["pdfs"]) if f.endswith(".pdf")]) == 1

    def test_an_empty_run_does_not_disturb_an_existing_index(self, tmp_path):
        root = str(tmp_path / "Lit")
        paths = index.ensure_layout(root)
        stage(paths["ingest"], "a.pdf", "Coral reef cover through time", "10.1/a")
        ingest.run(root, use_network=False)
        before = open(paths["index"], "rb").read()

        ingest.run(root, use_network=False)

        assert open(paths["index"], "rb").read() == before


class TestContractsThatMustNotDrift:
    def test_the_column_set_is_exactly_what_the_lab_was_promised(self):
        assert index.COLUMNS == [
            "link", "key", "filename", "authors", "first_author", "year", "title", "journal",
            "doi", "url", "citations", "citations_retrieved", "citations_per_year",
            "impact", "tags_all", "topic_tags", "method_tags", "region_tags", "taxa_tags",
            "vicar_relevance", "summary", "key_findings", "source", "date_added", "notes",
        ]

    def test_link_is_the_first_column_so_a_reader_can_click_straight_through(self):
        assert index.COLUMNS[0] == "link"

    def test_tags_all_sits_next_to_the_facets_it_summarizes(self):
        assert index.COLUMNS[index.COLUMNS.index("tags_all") + 1] == "topic_tags"

    def test_the_workbook_column_order_matches_the_index(self, tmp_path):
        from openpyxl import load_workbook

        out = sheet.build_workbook([{"key": "k"}], str(tmp_path / "w.xlsx"))
        headers = [c.value for c in load_workbook(out)[sheet.SHEET_TITLE][1]]
        # The workbook appends one computed column the CSV does not carry.
        assert headers[:-1] == index.COLUMNS
        assert headers[-1] == sheet.MATCH_COLUMN

    def test_every_column_has_a_declared_width(self):
        assert set(sheet.COLUMN_WIDTHS) >= set(index.COLUMNS)

    def test_data_rows_stay_short_enough_to_read_as_a_table(self):
        assert sheet.DATA_ROW_HEIGHT <= 22

    def test_the_filename_standard_is_author_year_shorttitle(self):
        key = naming.build_key("Nemeth, Richard S.", 2005, "Red hind spawning aggregation recovery")
        assert key == "Nemeth_2005_RedHindSpawningAggregationRecovery"
        assert naming.build_filename(key) == f"{key}.pdf"

    def test_a_key_is_always_recoverable_from_its_filename(self):
        key = "Nemeth_2005_RedHindSpawningAggregationRecovery"
        assert os.path.splitext(naming.build_filename(key))[0] == key


class TestIndexAndFolderStayInStep:
    def test_the_index_never_lists_a_paper_with_no_file(self, tmp_path):
        root = str(tmp_path / "Lit")
        paths = index.ensure_layout(root)
        for n, doi in enumerate(["10.1/a", "10.1/b", "10.1/c"]):
            stage(paths["ingest"], f"p{n}.pdf", f"Study number {n} of reef change", doi)
        ingest.run(root, use_network=False)

        rows = index.read_index(paths["index"])
        on_disk = set(os.listdir(paths["pdfs"]))
        assert {row["filename"] for row in rows} == on_disk

    def test_a_rejected_file_never_reaches_the_index_or_the_folder(self, tmp_path, broken_files):
        root = str(tmp_path / "Lit")
        paths = index.ensure_layout(root)
        shutil.copy(broken_files["empty"], os.path.join(paths["ingest"], "broken.pdf"))
        result = ingest.run(root, use_network=False)

        assert result["rejected"] == 1
        assert index.read_index(paths["index"]) == []
        assert [f for f in os.listdir(paths["pdfs"]) if f.endswith(".pdf")] == []

    def test_a_rejected_file_stays_in_ingest_for_the_operator_to_see(self, tmp_path, broken_files):
        root = str(tmp_path / "Lit")
        paths = index.ensure_layout(root)
        shutil.copy(broken_files["empty"], os.path.join(paths["ingest"], "broken.pdf"))
        ingest.run(root, use_network=False)

        assert "broken.pdf" in os.listdir(paths["ingest"])


class TestAuthorAbbreviationInTheSheet:
    """The CSV keeps every author. The sheet shows a readable form."""

    def test_a_short_list_is_untouched(self):
        authors = "Smith, Tyler B.; Brandt, Marilyn E."
        assert sheet.abbreviate_authors(authors) == authors

    def test_a_list_at_the_limit_is_untouched(self):
        authors = "; ".join(f"Author{n}" for n in range(sheet.MAX_AUTHORS_SHOWN))
        assert sheet.abbreviate_authors(authors) == authors

    def test_a_long_list_is_shortened_and_counted(self):
        authors = "; ".join(f"Author{n}" for n in range(290))
        result = sheet.abbreviate_authors(authors)
        assert result.count(";") == sheet.MAX_AUTHORS_SHOWN
        assert "et al. (290 authors)" in result
        assert len(result) < 200

    @pytest.mark.parametrize("empty", [None, "", "   "])
    def test_empty_input_is_safe(self, empty):
        assert sheet.abbreviate_authors(empty) == ""

    def test_the_csv_still_holds_every_author(self, tmp_path):
        full = "; ".join(f"Author{n}" for n in range(290))
        path = str(tmp_path / index.INDEX_FILENAME)
        index.write_index(path, [{"key": "k", "authors": full}])
        assert index.read_index(path)[0]["authors"] == full

    def test_the_workbook_shows_the_short_form(self, tmp_path):
        from openpyxl import load_workbook

        full = "; ".join(f"Author{n}" for n in range(290))
        out = sheet.build_workbook([{"key": "k", "authors": full}], str(tmp_path / "w.xlsx"))
        loaded = load_workbook(out)[sheet.SHEET_TITLE]
        header = [c.value for c in loaded[1]]
        values = {header[i]: c.value for i, c in enumerate(loaded[2])}
        assert "et al." in values["authors"]
        assert len(values["authors"]) < 200
