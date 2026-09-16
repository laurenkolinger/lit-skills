"""Build the formatted workbook that becomes the VICAR Lit Google Sheet.

Every long text column clips rather than wraps, and every data row keeps a fixed short
height, so the sheet stays compact. A reader clicks a cell to read the full summary.
"""

import os

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from .index import COLUMNS, TAG_FACETS, build_tags_all

SHEET_TITLE = "Lit index"
SEARCH_TITLE = "Search"
TAGS_TITLE = "Tags"
LINK_TEXT = "open"
DATA_ROW_HEIGHT = 20
HEADER_ROW_HEIGHT = 26
HEADER_FILL = "1F4E5A"
HEADER_FONT_COLOR = "FFFFFF"
DEFAULT_COLUMN_WIDTH = 16

# Column widths chosen so the columns a reader scans stay legible and the long text columns
# stay narrow enough that the sheet does not sprawl sideways.
COLUMN_WIDTHS = {
    "link": 7,
    "key": 34,
    "filename": 34,
    "authors": 30,
    "first_author": 16,
    "year": 7,
    "title": 46,
    "journal": 24,
    "doi": 24,
    "url": 24,
    "citations": 10,
    "citations_retrieved": 13,
    "citations_per_year": 12,
    "impact": 13,
    "tags_all": 40,
    "topic_tags": 26,
    "method_tags": 26,
    "region_tags": 16,
    "taxa_tags": 18,
    "vicar_relevance": 26,
    "summary": 60,
    "key_findings": 60,
    "source": 18,
    "date_added": 12,
    "notes": 24,
}

NUMERIC_COLUMNS = {"citations", "citations_per_year", "year"}

# A 290-author consortium report makes one cell 5,000 characters wide. The CSV keeps every
# name; the sheet shows a readable form, because the sheet is the reading surface.
MAX_AUTHORS_SHOWN = 8


def abbreviate_authors(value, limit=MAX_AUTHORS_SHOWN):
    """Shorten a long author list for display, leaving short ones untouched.

    Parameters:
        value (str | None): the full author list, semicolon separated.
        limit (int): how many authors to show before adding "et al.".

    Returns:
        str: the display form of the author list.
    """
    if not value or not str(value).strip():
        return ""
    names = [n.strip() for n in str(value).split(";") if n.strip()]
    if len(names) <= limit:
        return str(value)
    return "; ".join(names[:limit]) + f"; et al. ({len(names)} authors)"


def _coerce_number(column, value):
    """Convert a numeric column's text back to a number so the sheet can sort and filter it.

    Parameters:
        column (str): the column name.
        value: the cell value as stored in the CSV.

    Returns:
        int | float | str: a number when the column is numeric and the value parses,
        otherwise the value unchanged.
    """
    if column not in NUMERIC_COLUMNS or value in (None, ""):
        return value
    text = str(value).strip()
    try:
        return int(text)
    except ValueError:
        pass
    try:
        return float(text)
    except ValueError:
        return value


def build_workbook(rows, output_path, columns=None):
    """Write the index rows to a formatted xlsx file.

    Parameters:
        rows (list[dict]): index rows.
        output_path (str): where to write the workbook.
        columns (list[str] | None): column order, defaulting to the index columns.

    Returns:
        str: the path written.

    Raises:
        ValueError: when ``output_path`` is empty.
    """
    if not output_path:
        raise ValueError("build_workbook needs an output path")
    headers = list(columns or COLUMNS)

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = SHEET_TITLE

    header_font = Font(bold=True, color=HEADER_FONT_COLOR, size=11)
    header_fill = PatternFill("solid", fgColor=HEADER_FILL)
    header_alignment = Alignment(horizontal="left", vertical="center", wrap_text=False)
    # wrap_text stays False everywhere: that is what keeps rows short and makes text clip.
    data_alignment = Alignment(horizontal="left", vertical="center", wrap_text=False)

    for position, name in enumerate(headers, start=1):
        cell = sheet.cell(row=1, column=position, value=name)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_alignment
        letter = get_column_letter(position)
        sheet.column_dimensions[letter].width = COLUMN_WIDTHS.get(name, DEFAULT_COLUMN_WIDTH)
    sheet.row_dimensions[1].height = HEADER_ROW_HEIGHT

    for offset, row in enumerate(rows):
        row = dict(row)
        row["tags_all"] = build_tags_all(row)
        row_number = offset + 2
        for position, name in enumerate(headers, start=1):
            value = row.get(name, "")
            if name == "authors":
                value = abbreviate_authors(value)
            if name == "link" and value:
                cell = sheet.cell(row=row_number, column=position, value=LINK_TEXT)
                cell.hyperlink = str(value)
                cell.font = Font(color="0563C1", underline="single", size=11)
            else:
                cell = sheet.cell(row=row_number, column=position, value=_coerce_number(name, value))
            cell.alignment = data_alignment
        sheet.row_dimensions[row_number].height = DATA_ROW_HEIGHT

    last_column = get_column_letter(len(headers))
    last_row = len(rows) + 1
    sheet.freeze_panes = "C2"
    sheet.auto_filter.ref = f"A1:{last_column}{last_row}"
    sheet.sheet_view.showGridLines = True

    _add_search_tab(workbook, headers, len(rows))
    _add_tags_tab(workbook, rows)

    parent = os.path.dirname(os.path.abspath(output_path))
    if parent:
        os.makedirs(parent, exist_ok=True)
    workbook.save(output_path)
    return output_path


def _column_letter(headers, name):
    """Find the spreadsheet column letter for a named index column.

    Parameters:
        headers (list[str]): the column order written to the sheet.
        name (str): the column name.

    Returns:
        str: the column letter, or ``A`` when the column is absent.
    """
    return get_column_letter(headers.index(name) + 1) if name in headers else "A"


def _add_search_tab(workbook, headers, row_count):
    """Add a tab that filters the index by up to three tags plus free text.

    The formulas are Google Sheets formulas. Each tag is matched with its surrounding pipes,
    so "AUV" matches the AUV tag and not a tag that merely contains those letters. An empty
    input cell drops out of the filter rather than matching nothing.

    Parameters:
        workbook (openpyxl.Workbook): the workbook being built.
        headers (list[str]): the column order on the index sheet.
        row_count (int): how many data rows the index sheet holds.

    Returns:
        None
    """
    sheet = workbook.create_sheet(SEARCH_TITLE, 0)
    last = row_count + 1
    tags = _column_letter(headers, "tags_all")
    quoted = f"'{SHEET_TITLE}'"
    columns = {name: _column_letter(headers, name) for name in
               ("link", "key", "year", "title", "impact", "citations", "tags_all", "summary")}

    sheet["A1"] = "Search the VICAR lit library"
    sheet["A1"].font = Font(bold=True, size=14)
    sheet["A3"] = "Tag 1"
    sheet["A4"] = "Tag 2"
    sheet["A5"] = "Tag 3"
    sheet["A6"] = "Words in title or summary"
    for cell in ("A3", "A4", "A5", "A6"):
        sheet[cell].font = Font(bold=True)
    sheet["C3"] = "Type a tag, for example AUV. Leave a box empty to ignore it."
    sheet["C4"] = "Tags must match the Tags tab exactly. All boxes are combined with AND."
    sheet["C6"] = "Free text, for example bleaching or St. Thomas."
    for cell in ("C3", "C4", "C6"):
        sheet[cell].font = Font(italic=True, color="666666")

    for cell in ("B3", "B4", "B5", "B6"):
        sheet[cell].fill = PatternFill("solid", fgColor="FFF2CC")
        sheet[cell].border = Border(*[Side(style="thin", color="BFBFBF")] * 4)

    sheet["A8"] = "Matching papers"
    sheet["A8"].font = Font(bold=True)
    for offset, label in enumerate(["link", "key", "year", "title", "impact", "citations", "tags"]):
        cell = sheet.cell(row=9, column=offset + 1, value=label)
        cell.font = Font(bold=True, color=HEADER_FONT_COLOR)
        cell.fill = PatternFill("solid", fgColor=HEADER_FILL)

    def tag_condition(box):
        """Build the FILTER condition for one tag input box."""
        return (f'IF(${box}="",TRUE,ISNUMBER(SEARCH("|"&${box}&"|",'
                f'{quoted}!${tags}$2:${tags}${last})))')

    title_col, summary_col = columns["title"], columns["summary"]
    text_condition = (
        f'IF($B$6="",TRUE,'
        f'ISNUMBER(SEARCH($B$6,{quoted}!${title_col}$2:${title_col}${last}))+'
        f'ISNUMBER(SEARCH($B$6,{quoted}!${summary_col}$2:${summary_col}${last})))'
    )
    shown = ", ".join(
        f'{quoted}!${columns[name]}$2:${columns[name]}${last}'
        for name in ("link", "key", "year", "title", "impact", "citations", "tags_all")
    )
    sheet["A10"] = (
        f'=IFERROR(SORT(FILTER({{{shown}}}, '
        f'{tag_condition("B$3")}, {tag_condition("B$4")}, {tag_condition("B$5")}, '
        f'{text_condition}), 6, FALSE), "No papers match. Check the spelling against the Tags tab.")'
    )

    widths = {"A": 7, "B": 36, "C": 7, "D": 58, "E": 14, "F": 11, "G": 46}
    for letter, width in widths.items():
        sheet.column_dimensions[letter].width = width
    sheet.freeze_panes = "A10"


def _add_tags_tab(workbook, rows):
    """Add a tab listing every tag in use, with how many papers carry it.

    Parameters:
        workbook (openpyxl.Workbook): the workbook being built.
        rows (list[dict]): the index rows.

    Returns:
        None
    """
    sheet = workbook.create_sheet(TAGS_TITLE)
    counts = {}
    for row in rows:
        for facet in TAG_FACETS:
            for tag in str(row.get(facet) or "").split(","):
                tag = tag.strip()
                if tag:
                    counts.setdefault(facet, {})
                    counts[facet][tag] = counts[facet].get(tag, 0) + 1

    sheet["A1"] = "Every tag in the library, and how many papers carry it"
    sheet["A1"].font = Font(bold=True, size=14)
    sheet["A2"] = "Copy a tag into the Search tab to filter by it."
    sheet["A2"].font = Font(italic=True, color="666666")

    for position, label in enumerate(["facet", "tag", "papers"]):
        cell = sheet.cell(row=4, column=position + 1, value=label)
        cell.font = Font(bold=True, color=HEADER_FONT_COLOR)
        cell.fill = PatternFill("solid", fgColor=HEADER_FILL)

    line = 5
    for facet in TAG_FACETS:
        for tag, count in sorted(counts.get(facet, {}).items(), key=lambda kv: (-kv[1], kv[0])):
            sheet.cell(row=line, column=1, value=facet.replace("_tags", "").replace("_", " "))
            sheet.cell(row=line, column=2, value=tag)
            sheet.cell(row=line, column=3, value=count)
            sheet.row_dimensions[line].height = DATA_ROW_HEIGHT
            line += 1

    sheet.column_dimensions["A"].width = 18
    sheet.column_dimensions["B"].width = 34
    sheet.column_dimensions["C"].width = 10
    sheet.freeze_panes = "A5"
    if line > 5:
        sheet.auto_filter.ref = f"A4:C{line - 1}"
